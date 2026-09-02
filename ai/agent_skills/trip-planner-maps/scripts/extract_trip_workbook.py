#!/usr/bin/env python3

import argparse
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path


try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - environment guard
    raise SystemExit(
        "openpyxl is required for extract_trip_workbook.py. Install it with `python3 -m pip install openpyxl`."
    ) from exc


SPACE_RE = re.compile(r"\s+")
CHECKBOX_RE = re.compile(r"^\[\s*[xX ]\s*\]\s*")
NUMBERING_RE = re.compile(r"^\d+[\.\)]\s*")
BULLET_RE = re.compile(r"^[-*]\s*")


def normalize_text(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = str(value).replace("\xa0", " ").strip()
    text = SPACE_RE.sub(" ", text)
    return text or None


def normalized_cells(row):
    return [normalize_text(cell) for cell in row]


def first_nonempty(cells):
    for cell in cells:
        if cell:
            return cell
    return None


def clean_stop_candidate(text):
    if not text:
        return None

    cleaned = CHECKBOX_RE.sub("", text)
    cleaned = NUMBERING_RE.sub("", cleaned)
    cleaned = BULLET_RE.sub("", cleaned)
    cleaned = cleaned.strip(" :")
    cleaned = SPACE_RE.sub(" ", cleaned)
    return cleaned or None


def looks_like_heading(text):
    if not text:
        return False

    stripped = text.strip()
    if stripped.isupper() and len(stripped.split()) <= 6:
        return True

    lowered = stripped.lower()
    heading_prefixes = (
        "top ",
        "for your convenience",
        "downloadable resources",
        "maps & guides",
        "activities & dining",
        "culture & history",
        "neighborhoods & parks",
        "unique experiences",
        "markets",
        "iconic landmarks",
        "world-class free museums",
        "palaces & imperial history",
    )
    return lowered.startswith(heading_prefixes)


def looks_like_stop(text):
    if not text or looks_like_heading(text):
        return False

    lowered = text.lower()
    if lowered.startswith(("http://", "https://", "www.")):
        return False
    if len(text.split()) > 20:
        return False
    return True


def infer_location_matches(text, location_names):
    lowered = text.lower()
    return [name for name in location_names if name.lower() in lowered]


def parse_overview(ws, location_names):
    rows = []
    matched_locations = []

    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = normalized_cells(row)
        if not any(cells):
            continue
        if row_index == 1 and cells[:4] == ["Date", "Day", "Location", "Plans"]:
            continue

        date_value = cells[0] if len(cells) > 0 else None
        day_value = cells[1] if len(cells) > 1 else None
        location_value = cells[2] if len(cells) > 2 else None
        plans_value = cells[3] if len(cells) > 3 else None
        location_matches = infer_location_matches(location_value or "", location_names)
        matched_locations.extend(location_matches)

        rows.append(
            {
                "date": date_value,
                "day": day_value,
                "location": location_value,
                "plans": plans_value,
                "location_matches": location_matches,
            }
        )

    deduped_sequence = []
    for name in matched_locations:
        if not deduped_sequence or deduped_sequence[-1] != name:
            deduped_sequence.append(name)

    return {
        "rows": rows,
        "city_sequence": deduped_sequence,
    }


def parse_city_sheet(ws):
    raw_rows = []
    stop_candidates = []

    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = normalized_cells(row)
        if not any(cells):
            continue

        first_cell = first_nonempty(cells)
        stop_name = clean_stop_candidate(first_cell)
        row_kind = "note"
        if looks_like_heading(first_cell):
            row_kind = "heading"
        elif looks_like_stop(stop_name):
            row_kind = "candidate_stop"

        raw_rows.append(
            {
                "row_index": row_index,
                "cells": [cell for cell in cells if cell],
                "row_kind": row_kind,
            }
        )

        if row_kind == "candidate_stop" and stop_name:
            stop_candidates.append(stop_name)

    return {
        "rows": raw_rows,
        "stop_candidates": stop_candidates,
    }


def analyze_missing_inputs(full_text, overview_rows):
    missing = []

    if not overview_rows:
        missing.append("trip dates and city sequence")
    elif not any(row.get("date") for row in overview_rows):
        missing.append("trip dates")

    hotel_keywords = (
        "hotel",
        "hostel",
        "airbnb",
        "accommodation",
        "stay at",
        "lodging",
        "check-in",
        "check in",
    )
    if not any(keyword in text for text in full_text for keyword in hotel_keywords):
        missing.append("lodging or hotel base per city")

    hard_timing_keywords = ("am", "pm", "flight", "train", "reservation", "ticket", "check-in", "check in")
    if not any(keyword in text for text in full_text for keyword in hard_timing_keywords):
        missing.append("fixed booking times or arrival/departure times")

    return missing


def workbook_summary(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_names = wb.sheetnames
    location_names = [name for name in sheet_names if name.lower() != "overview"]

    summary = {
        "workbook_path": str(path.resolve()),
        "workbook_name": path.name,
        "sheet_names": sheet_names,
        "overview": None,
        "city_sheets": {},
        "planning_signals": {},
    }

    all_text = []
    city_name_groups = defaultdict(list)

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        if sheet_name.lower() == "overview":
            overview = parse_overview(ws, location_names)
            summary["overview"] = overview
            for row in overview["rows"]:
                all_text.extend(filter(None, (row.get("location"), row.get("plans"))))
            continue

        parsed = parse_city_sheet(ws)
        summary["city_sheets"][sheet_name] = parsed
        city_name_groups[sheet_name.split()[0].lower()].append(sheet_name)

        for row in parsed["rows"]:
            all_text.extend(row["cells"])

    overview_rows = summary["overview"]["rows"] if summary["overview"] else []
    missing_inputs = analyze_missing_inputs([text.lower() for text in all_text], overview_rows)

    summary["planning_signals"] = {
        "has_overview_sheet": summary["overview"] is not None,
        "has_dates": any(row.get("date") for row in overview_rows),
        "duplicate_city_groups": {
            key: names for key, names in city_name_groups.items() if len(names) > 1
        },
        "missing_inputs": missing_inputs,
    }

    return summary


def main():
    parser = argparse.ArgumentParser(
        description="Extract a trip workbook into structured JSON for itinerary planning."
    )
    parser.add_argument("workbook", help="Path to the .xlsx workbook")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    summary = workbook_summary(workbook_path)
    json.dump(summary, sys.stdout, indent=2)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
